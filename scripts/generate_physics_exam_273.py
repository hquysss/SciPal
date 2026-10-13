"""Excel import workbook for the Vĩnh Long grade-11 physics mid-term, code 273.

Run: python scripts/generate_physics_exam_273.py
Writes data/de-vat-li-11-giua-ki-1-ma-273.xlsx for the admin import page. Phần IV (essay) is not
importable: the system grades mc, truefalse and short only. Answers follow the school's answer key.
"""
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parents[1] / "data" / "de-vat-li-11-giua-ki-1-ma-273.xlsx"
CODE = "vatly11-gk1-273"
SUBJ = "physics"
SOURCE = "Đề kiểm tra giữa kì I 2025–2026, Vật lí 11 (THPT chuyên Nguyễn Thiện Thành, mã đề 273)"

# (key, stem, [(id, text)], correct id, explanation)
MC = [
    ("c01", r"Phương trình li độ của một vật dao động điều hoà có dạng $x = A\cos(\omega t + \varphi)$. Phương trình gia tốc của vật là",
     [("A", r"$a = -A\omega^2\sin(\omega t + \varphi)$."), ("B", r"$a = A\omega^2\sin(\omega t + \varphi)$."),
      ("C", r"$a = A\omega^2\cos(\omega t + \varphi)$."), ("D", r"$a = -A\omega^2\cos(\omega t + \varphi)$.")], "D",
     r"Gia tốc là đạo hàm bậc hai của li độ: $a = -\omega^2 x = -A\omega^2\cos(\omega t + \varphi)$."),
    ("c02", "Đồ thị li độ theo thời gian của dao động điều hòa là một",
     [("A", "đường thẳng"), ("B", "đoạn thẳng"), ("C", "đường hình sin"), ("D", "đường tròn.")], "C",
     "Li độ biến thiên theo hàm cosin của thời gian nên đồ thị là đường hình sin."),
    ("c03", "Chọn câu sai: Năng lượng của một vật dao động điều hòa",
     [("A", "luôn luôn là một hằng số."), ("B", "biến thiên tuần hoàn theo thời gian với chu kì $T$."),
      ("C", "bằng thế năng của vật khi qua vị trí biên."), ("D", "bằng động năng của vật khi qua vị trí cân bằng.")], "B",
     "Cơ năng của vật dao động điều hoà không đổi theo thời gian; chỉ động năng và thế năng biến thiên."),
    ("c04", r"Một vật dao động điều hòa theo phương trình $x = 4\cos\left(2\pi t + \frac{\pi}{3}\right)$ cm. Pha ban đầu của vật là",
     [("A", r"$\varphi = -\frac{2\pi}{3}$ rad"), ("B", r"$\varphi = -\frac{\pi}{6}$ rad"),
      ("C", r"$\varphi = \frac{\pi}{3}$ rad"), ("D", r"$\varphi = \frac{5\pi}{6}$ rad")], "C",
     r"So với $x = A\cos(\omega t + \varphi)$ thì $\varphi = \frac{\pi}{3}$ rad."),
    ("c05", "Biên độ của dao động cơ tắt dần",
     [("A", "biến thiên điều hòa theo thời gian."), ("B", "giảm dần theo thời gian."),
      ("C", "không đổi theo thời gian."), ("D", "tăng dần theo thời gian.")], "B",
     "Lực cản làm cơ năng giảm nên biên độ giảm dần."),
    ("c06", r"Một chất điểm có khối lượng $m$ đang dao động điều hòa. Khi chất điểm có vận tốc $v$ thì động năng của nó là",
     [("A", r"$mv^2$."), ("B", r"$vm^2$."), ("C", r"$\frac{vm^2}{2}$."), ("D", r"$\frac{mv^2}{2}$.")], "D",
     r"Động năng $W_đ = \frac{1}{2}mv^2$."),
    ("c07", r"Một con lắc lò xo gồm vật nhỏ có khối lượng $m$ gắn vào một lò xo có độ cứng 50 N/m. Tác dụng lên vật ngoại lực cưỡng bức $F = 40\cos(10\pi t - \pi)$ (N) dọc theo trục lò xo thì hiện tượng cộng hưởng xảy ra. Lấy $\pi^2 = 10$. Giá trị của $m$ là",
     [("A", "0,05 g"), ("B", "0,05 kg"), ("C", "5 kg"), ("D", "5 g")], "B",
     r"Cộng hưởng: $\omega_0 = \omega = 10\pi$ rad/s, nên $m = \frac{k}{\omega^2} = \frac{50}{100\pi^2} = 0{,}05$ kg."),
    ("c08", r"Phương trình li độ của một vật dao động điều hoà có dạng $x = A\cos(\omega t + \varphi)$. Phương trình vận tốc của vật là",
     [("A", r"$v = -\omega A\sin(\omega t + \varphi)$."), ("B", r"$v = -\omega A\cos(\omega t + \varphi)$."),
      ("C", r"$v = \omega A\sin(\omega t + \varphi)$."), ("D", r"$v = \omega A\cos(\omega t + \varphi)$.")], "A",
     r"Vận tốc là đạo hàm của li độ: $v = x' = -\omega A\sin(\omega t + \varphi)$."),
    ("c09", r"Một vật dao động điều hòa theo phương trình $x = 4\cos\left(4\pi t + \frac{\pi}{3}\right)$ cm, $t$ tính bằng giây. Thời gian vật thực hiện được một dao động toàn phần là",
     [("A", "4 s."), ("B", "2 s."), ("C", "0,5 s."), ("D", "1 s.")], "C",
     r"$T = \frac{2\pi}{\omega} = \frac{2\pi}{4\pi} = 0{,}5$ s."),
    ("c10", "Tần số góc có đơn vị là", [("A", "rad/s."), ("B", "rad."), ("C", "cm."), ("D", "Hz.")], "A",
     "Tần số góc đo bằng rad/s."),
    ("c11", r"Trong phương trình dao động điều hòa $x = A\cos(\omega t + \varphi)$, đại lượng $(\omega t + \varphi)$ gọi là",
     [("A", "chu kì của dao động."), ("B", "pha của dao động."), ("C", "tần số góc của dao động."), ("D", "biên độ của dao động.")], "B",
     r"$(\omega t + \varphi)$ là pha dao động tại thời điểm $t$."),
    ("c12", r"Một vật dao động điều hoà có phương trình $x = 4\cos(20\pi t - \pi/6)$ cm. Tần số và pha ban đầu của dao động lần lượt là",
     [("A", r"10 Hz và $-\pi/6$ rad"), ("B", r"10 Hz và $\pi/6$ rad"), ("C", r"1/10 Hz và $-\pi/6$ rad"), ("D", r"1/10 Hz và $\pi/6$ rad")], "A",
     r"$f = \frac{\omega}{2\pi} = 10$ Hz, $\varphi = -\frac{\pi}{6}$ rad."),
]

# (key, stem, [(id, text, correct)], explanation)
TF = [
    ("t01", r"Một vật có khối lượng 100 g dao động theo phương trình $x = 4\cos(3\pi t + 0{,}5\pi)$ (cm).",
     [("a", r"Pha ban đầu của vật là $0{,}5\pi$ rad/s.", False), ("b", "Chiều dài quỹ đạo dao động của vật là 8 cm.", True),
      ("c", "Chu kì của vật trong quá trình dao động là 3 s.", False),
      ("d", r"Phương trình vận tốc của vật là $v = -12\pi\sin(3\pi t + 0{,}5\pi)$ cm/s.", True)],
     r"a) Sai: pha ban đầu là $0{,}5\pi$ rad (không phải rad/s). b) Đúng: $L = 2A = 8$ cm. c) Sai: $T = \frac{2\pi}{3\pi} = \frac{2}{3}$ s. d) Đúng: $v = -\omega A\sin(\omega t + \varphi)$."),
    ("t02", r"Cho đồ thị li độ theo thời gian của một con lắc lò xo, biết vật có khối lượng $m = 150$ g, lấy $\pi^2 = 10$. Đồ thị là đường hình sin đi qua gốc toạ độ theo chiều dương, đạt cực đại $+10$ cm tại $t = 0{,}5$ s và trở về vị trí cân bằng tại $t = 1$ s.",
     [("a", "Biên độ dao động của vật bằng 10 cm.", True),
      ("b", r"Phương trình vận tốc của vật là $v = -10\pi\sin\left(\pi t - \frac{\pi}{2}\right)$ cm/s.", True),
      ("c", "Động năng và thế năng của con lắc biến thiên tuần hoàn với chu kì 2 s.", False),
      ("d", r"Cơ năng của con lắc là $7{,}5\cdot 10^{-3}$ J.", True)],
     r"Chu kì dao động $T = 2$ s nên $\omega = \pi$ rad/s, $x = 10\cos\left(\pi t - \frac{\pi}{2}\right)$ cm. c) Sai: động năng, thế năng biến thiên với chu kì $\frac{T}{2} = 1$ s. d) $W = \frac{1}{2}m\omega^2A^2 = \frac{1}{2}\cdot 0{,}15\cdot 10\cdot 0{,}1^2 = 7{,}5\cdot 10^{-3}$ J."),
]

# (key, stem, answer, explanation)
SHORT = [
    ("s01", "Một vật dao động điều hòa có chu kì 2 s, biên độ 10 cm. Chiều dài quỹ đạo của vật là bao nhiêu mét?", "0.2",
     r"$L = 2A = 20$ cm $= 0{,}2$ m."),
    ("s02", r"Một vật dao động điều hòa theo phương trình $x = 3\cos(\pi t + \pi/2)$ cm. Pha dao động tại thời điểm $t = 1{,}5$ s là bao nhiêu $\pi$ (rad)?", "2",
     r"Pha $= \pi\cdot 1{,}5 + \frac{\pi}{2} = 2\pi$ rad."),
    ("s03", r"Một vật dao động điều hòa với biên độ 5 cm. Khi vật có li độ 2,5 cm thì tốc độ của vật là $5\sqrt{3}$ cm/s. Vận tốc cực đại của dao động là bao nhiêu cm/s?", "10",
     r"$v = \omega\sqrt{A^2 - x^2}$ cho $\omega = 2$ rad/s, nên $v_{max} = \omega A = 10$ cm/s."),
    ("s04", r"Một con lắc lò xo gồm lò xo nhẹ và vật nhỏ dao động điều hòa theo phương ngang với tần số góc 10 rad/s. Khi động năng và thế năng (mốc ở vị trí cân bằng) bằng nhau thì vận tốc của vật có độ lớn bằng $0{,}6\sqrt{2}$ m/s. Biên độ dao động của con lắc là bao nhiêu mét? (làm tròn đến chữ số hàng phần trăm)", "0.12",
     r"$W_đ = W_t \Rightarrow |x| = \frac{A}{\sqrt{2}}$, $|v| = \frac{\omega A}{\sqrt{2}} \Rightarrow A = \frac{0{,}6\sqrt{2}\cdot\sqrt{2}}{10} = 0{,}12$ m."),
]


def build() -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    q = wb.create_sheet("Questions")
    q.append(["question_key", "subject_slug", "type", "difficulty", "stem_vi", "stem_en", "answer_key", "explanation_vi",
              "explanation_en", "rubric_vi", "rubric_en", "_template_example", "source"])
    it = wb.create_sheet("Question items")
    it.append(["subject_slug", "question_key", "item_id", "text_vi", "text_en", "correct", "_template_example"])
    ex = wb.create_sheet("Exams")
    ex.append(["exam_code", "subject_slug", "grade", "title_vi", "title_en", "duration_minutes", "format", "year", "_template_example"])
    sec = wb.create_sheet("Exam sections")
    sec.append(["exam_code", "type", "difficulty", "count", "section_key", "title_vi", "title_en", "max_points", "_template_example"])

    def key(k: str) -> str:
        return f"{CODE}-{k}"

    for k, stem, options, answer, expl in MC:
        q.append([key(k), SUBJ, "mc", 1, stem, "", "", expl, "", "", "", "FALSE", SOURCE])
        for oid, text in options:
            it.append([SUBJ, key(k), oid, text, "", oid == answer, "FALSE"])
    for k, stem, items, expl in TF:
        q.append([key(k), SUBJ, "truefalse", 2, stem, "", "", expl, "", "", "", "FALSE", SOURCE])
        for iid, text, ok in items:
            it.append([SUBJ, key(k), iid, text, "", ok, "FALSE"])
    for k, stem, answer, expl in SHORT:
        q.append([key(k), SUBJ, "short", 3, stem, "", answer, expl, "", "Chấp nhận dấu phẩy hoặc dấu chấm thập phân.", "", "FALSE", SOURCE])

    ex.append([CODE, SUBJ, 11, "Kiểm tra giữa kì I Vật lí 11 – Mã đề 273 (Phần I–III)",
               "Grade 11 Physics mid-term I – Code 273 (Parts I–III)", 45, "generic", 2025, "FALSE"])
    sec.append([CODE, "mc", 1, len(MC), "", "", "", "", "FALSE"])
    sec.append([CODE, "truefalse", 2, len(TF), "", "", "", "", "FALSE"])
    sec.append([CODE, "short", 3, len(SHORT), "", "", "", "", "FALSE"])
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
