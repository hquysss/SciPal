# -*- coding: utf-8 -*-
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1,<4"]
# ///
# How to run: python generate_lesson_21_tin.py
from __future__ import annotations

import json
import os
import xml.sax.saxutils
import zipfile

import openpyxl
from openpyxl.styles import Font, PatternFill

OUT_DIR = r"d:\Code\SciPal\export-lesson"
TEMPLATE = r"d:\Code\SciPal\frontend\public\templates\scipal-lesson-template.docx"
SUBJECT = "informatics"
GRADE = 11
TOPIC_VI = "Kĩ thuật lập trình"
TOPIC_EN = "Programming techniques"
BOOK_VI = "Tin học 11 - Kết nối tri thức"
ORIENTATION_VI = "Khoa học máy tính"


def configure(grade, topic_vi, topic_en, orientation_vi=None):
    """Switch the shared writer from Tin 11 (default) to another grade/topic before build()."""
    global GRADE, TOPIC_VI, TOPIC_EN, BOOK_VI, ORIENTATION_VI
    GRADE, TOPIC_VI, TOPIC_EN = grade, topic_vi, topic_en
    BOOK_VI = f"Tin học {grade} - Kết nối tri thức"
    ORIENTATION_VI = orientation_vi


def trace_block(copy, source, inputs=()):
    """A `code-trace` block: run the Python once now and keep the recorded steps."""
    import subprocess
    import sys
    import tempfile

    code_text = source.strip("\n") + "\n"
    # A child process via the CLI: multiprocessing on Windows would re-import this lesson script.
    with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False, encoding="utf-8") as handle:
        handle.write(code_text)
    args = [sys.executable, "-X", "utf8", os.path.join(os.path.dirname(os.path.abspath(__file__)), "code_trace.py"), handle.name]
    for value in inputs:
        args += ["--input", value]
    run = subprocess.run(args, capture_output=True, encoding="utf-8", check=False)
    os.unlink(handle.name)
    if run.returncode != 0:
        raise RuntimeError(f"code_trace failed: {run.stderr.strip() or run.stdout.strip()}")
    result = json.loads(run.stdout)
    config = {"code": code_text, "input": list(inputs), "steps": result["steps"]}
    if result.get("truncated"):
        config["truncated"] = True
    return interactive("code-trace", copy, config)


def web_block(copy, html, css="", height=360):
    """A `web-preview` block: HTML body + CSS shown line by line."""
    return interactive("web-preview", copy, {"html": html.strip("\n"), "css": css.strip("\n"), "height": height})


def T(vi, en):
    return {"type": "theory", "content": {"vi": vi, "en": en}}


def code(python, cpp):
    return {
        "type": "code",
        "tabs": [
            {"lang": "python", "code": python.strip()},
            {"lang": "cpp", "code": cpp.strip()},
        ],
    }


def interactive(kind, copy, config):
    return {
        "type": "interactive",
        "kind": kind,
        "heading": {"vi": copy["heading_vi"], "en": copy["heading_en"]},
        "caption": {"vi": copy["caption_vi"], "en": copy["caption_en"]},
        "offline": True,
        "config": config,
    }


def option(letter, vi, en, ok=False):
    return (letter, vi, en, ok)


def mc(key, level, q, q_en, why, why_en, right, *opts):
    """Multiple choice. `opts` are (vi, en) in order A, B, C, D; `right` is the letter of the correct one."""
    items = [option("ABCD"[i], vi, en, "ABCD"[i] == right) for i, (vi, en) in enumerate(opts)]
    return (key, level, q, q_en, why, why_en, items)


def tf(key, level, q, q_en, why, why_en, *stmts):
    """True/false cluster. `stmts` are (vi, en, is_true) for statements a, b, c, d."""
    items = [option("abcd"[i], vi, en, bool(ok)) for i, (vi, en, ok) in enumerate(stmts)]
    return (key, level, q, q_en, why, why_en, items)


def _lesson_text(lesson, include_quiz):
    content = [
        "SCIPAL-LESSON-V1",
        f"# Bài học: {BOOK_VI} - {lesson['title_vi']}",
        f"# Môn: informatics (Tin học) | Lớp: {GRADE}"
        + (f" | Định hướng: {ORIENTATION_VI}" if ORIENTATION_VI else ""),
        f"subject: {SUBJECT}",
        f"grade: {GRADE}",
        f"topic_vi: {TOPIC_VI}",
        f"topic_en: {TOPIC_EN}",
        f"title_vi: {lesson['title_vi']}",
        f"title_en: {lesson['title_en']}\n",
    ]
    for block in lesson["blocks"]:
        if block["type"] == "theory":
            content += [
                "[THEORY]",
                f"vi: {block['content']['vi']}",
                f"en: {block['content']['en']}\n",
            ]
        elif block["type"] == "code":
            for tab in block["tabs"]:
                content += [f"[CODE:{tab['lang']}]", f"{tab['code']}\n"]
        elif block["type"] == "interactive":
            content += [
                f"[INTERACTIVE:{block['kind']}]",
                f"heading_vi: {block['heading']['vi']}",
                f"heading_en: {block['heading']['en']}",
                f"caption_vi: {block['caption']['vi']}",
                f"caption_en: {block['caption']['en']}",
                f"config_json: {json.dumps(block['config'], ensure_ascii=False)}\n",
            ]
        else:
            raise ValueError(f"Unsupported block type: {block['type']}")
    if include_quiz:
        content.extend(f"[QUIZ:{question[0]}]" for question in lesson["questions"] + lesson["truefalse"])
    return "\n".join(content).rstrip() + "\n"


def _make_docx(text, out_path):
    lines = text.strip().replace("\r\n", "\n").split("\n")
    body = "".join(
        f'<w:p><w:r><w:t xml:space="preserve">{xml.sax.saxutils.escape(line)}</w:t></w:r></w:p>'
        for line in lines
    ) + "<w:sectPr/>"
    doc_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    with zipfile.ZipFile(TEMPLATE, "r") as source, zipfile.ZipFile(
        out_path, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            value = doc_xml.encode("utf-8") if item.filename == "word/document.xml" else source.read(item.filename)
            target.writestr(item, value)


def _setup_sheet(sheet, headers, widths):
    sheet.append(headers)
    header_fill = PatternFill(start_color="166534", end_color="166534", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for index, _ in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=index)
        cell.fill, cell.font = header_fill, header_font
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = widths[index - 1]
    sheet.freeze_panes = "A2"


def _make_workbook(lesson, out_path):
    workbook = openpyxl.Workbook()
    questions_sheet = workbook.active
    questions_sheet.title = "Questions"
    _setup_sheet(
        questions_sheet,
        [
            "question_key", "subject_slug", "type", "difficulty", "stem_vi", "stem_en",
            "answer_key", "explanation_vi", "explanation_en", "rubric_vi", "rubric_en",
        ],
        [32, 16, 12, 12, 50, 50, 15, 45, 45, 20, 20],
    )
    items_sheet = workbook.create_sheet(title="Question items")
    _setup_sheet(
        items_sheet,
        ["subject_slug", "question_key", "item_id", "text_vi", "text_en", "correct"],
        [16, 32, 10, 48, 48, 12],
    )
    for question_type, rows in (("mc", lesson["questions"]), ("truefalse", lesson["truefalse"])):
        for key, difficulty, stem_vi, stem_en, explanation_vi, explanation_en, items in rows:
            questions_sheet.append(
                (key, SUBJECT, question_type, difficulty, stem_vi, stem_en, "", explanation_vi, explanation_en)
            )
            for item_id, text_vi, text_en, correct in items:
                items_sheet.append((SUBJECT, key, item_id, text_vi, text_en, correct))
    _setup_sheet(
        workbook.create_sheet(title="Exams"),
        ["exam_code", "subject_slug", "grade", "title_vi", "title_en", "duration_minutes"],
        [15, 15, 10, 30, 30, 18],
    )
    _setup_sheet(
        workbook.create_sheet(title="Exam sections"),
        ["exam_code", "type", "difficulty", "count"],
        [15, 15, 15, 15],
    )
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise ValueError(f"{sheet.title}!{cell.coordinate} starts with '=': Excel reads it as a formula; reword it")
    workbook.save(out_path)


def build(lesson):
    os.makedirs(OUT_DIR, exist_ok=True)
    slug = lesson["slug"]
    with open(os.path.join(OUT_DIR, slug + ".json"), "w", encoding="utf-8") as output:
        json.dump(
            {"title_vi": lesson["title_vi"], "title_en": lesson["title_en"], "blocks": lesson["blocks"]},
            output,
            ensure_ascii=False,
            indent=2,
        )
    full_text = _lesson_text(lesson, include_quiz=True)
    lesson_text = _lesson_text(lesson, include_quiz=False)
    for name, text in ((slug + ".txt", full_text), (slug + "-khong-quiz.txt", lesson_text)):
        with open(os.path.join(OUT_DIR, name), "w", encoding="utf-8") as output:
            output.write(text)
    _make_docx(full_text, os.path.join(OUT_DIR, slug + ".docx"))
    _make_docx(lesson_text, os.path.join(OUT_DIR, slug + "-khong-quiz.docx"))
    _make_workbook(lesson, os.path.join(OUT_DIR, slug + "-cau-hoi.xlsx"))
    print(
        "Done:", slug, len(lesson["blocks"]), "blocks,",
        len(lesson["questions"]), "mc +", len(lesson["truefalse"]), "tf",
    )
